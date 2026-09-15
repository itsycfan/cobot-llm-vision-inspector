#!/usr/bin/env python3
"""
=============================================================================
Assembly Detection System - SIMPLIFIED VERSION 6.0
=============================================================================
Focus: 
- Clear code structure
- Optimized GENERIC prompt (result-oriented)
- Easy debugging and analysis
- Save all results and frames

Key Improvement:
- Emphasize OUTCOME not just PROCESS
- Judge differences by TASK COMPLETION not just visual appearance
- Use Success/Failure criteria from Excel

Version: 6.0 (Simplified Result-Oriented)
Date: 2025-10-28
=============================================================================
"""

import os
import json
import time
import base64
import cv2
import numpy as np
from datetime import datetime
from typing import Dict, List, Optional
from dataclasses import dataclass, field, asdict
from pathlib import Path

# Suppress warnings
os.environ['TF_CPP_MIN_LOG_LEVEL'] = '3'
import warnings
warnings.filterwarnings('ignore')

try:
    import mediapipe as mp
    MP_AVAILABLE = True
except:
    MP_AVAILABLE = False

try:
    import openai
    OPENAI_AVAILABLE = True
except:
    OPENAI_AVAILABLE = False

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except:
    EXCEL_AVAILABLE = False


# =============================================================================
# 1. CONFIGURATION
# =============================================================================

@dataclass
class Config:
    """System configuration"""
    
    # API
    OPENAI_API_KEY: str = os.getenv("OPENAI_API_KEY", "")
    GPT_MODEL: str = "gpt-4o"
    GPT_TEMPERATURE: float = 0.1
    
    # Paths
    VIDEO_FOLDER: str = "VideoExample"
    RESULTS_FOLDER: str = "Results_V6"
    EXCEL_PATH: str = "Data/AssemblyInstructions.xlsx"
    REFERENCE_VIDEO: str = "OK1"
    
    # Frame settings
    KEY_FRAMES: int = 8
    REF_FRAMES: int = 6
    IMAGE_WIDTH: int = 768
    JPEG_QUALITY: int = 85
    
    # Decision thresholds
    CRITICAL_THRESHOLD: int = 1
    MAJOR_THRESHOLD: int = 1
    SCORE_THRESHOLD: float = 70.0
    
    def __post_init__(self):
        """Create directories"""
        os.makedirs(self.RESULTS_FOLDER, exist_ok=True)
        os.makedirs(os.path.join(self.RESULTS_FOLDER, "frames"), exist_ok=True)
        os.makedirs(os.path.join(self.RESULTS_FOLDER, "reports"), exist_ok=True)


# Ground truth
GROUND_TRUTH = {
    "OK1": "OK", "OK2": "OK", "OK3": "OK", "OK4": "OK", "OK5": "OK",
    "KO1": "KO", "KO2": "KO", "KO3": "KO", "KO4": "KO", "KO5": "KO"
}


# =============================================================================
# 2. DATA MODELS
# =============================================================================

@dataclass
class Difference:
    """A single observed difference"""
    description: str
    is_problem: bool
    severity: str  # CRITICAL, MAJOR, MINOR, NEGLIGIBLE
    category: str  # safety, spatial, operation


@dataclass
class Result:
    """Analysis result"""
    video_name: str
    timestamp: str
    
    # Scores
    safety_score: float
    spatial_score: float
    operation_score: float
    overall_score: float
    
    # Analysis
    differences: List[Difference] = field(default_factory=list)
    summary: str = ""
    
    # Decision
    decision: str = ""  # PASS or FAIL
    reason: str = ""
    
    # Validation
    ground_truth: str = ""
    correct: bool = False
    
    @property
    def critical(self) -> List[Difference]:
        return [d for d in self.differences if d.is_problem and d.severity == 'CRITICAL']
    
    @property
    def major(self) -> List[Difference]:
        return [d for d in self.differences if d.is_problem and d.severity == 'MAJOR']
    
    @property
    def minor(self) -> List[Difference]:
        return [d for d in self.differences if d.is_problem and d.severity == 'MINOR']
    
    @property
    def negligible(self) -> List[Difference]:
        return [d for d in self.differences if not d.is_problem]


# =============================================================================
# 3. UTILITIES
# =============================================================================

def load_task_data(config: Config) -> Optional[Dict]:
    """Load task data from Excel"""
    
    if not EXCEL_AVAILABLE:
        return None
    
    try:
        wb = openpyxl.load_workbook(config.EXCEL_PATH)
        ws = wb.active
        
        # Get headers
        headers = [cell.value for cell in ws[1]]
        
        # Get data from row 2
        data = {}
        for idx, header in enumerate(headers, 1):
            if header:
                value = ws.cell(row=2, column=idx).value
                if value:
                    data[header] = value
        
        wb.close()
        
        return {
            'task_description': data.get('TaskDescription', ''),
            'safety_requirements': data.get('SafetyRequirements', ''),
            'spatial_requirements': data.get('SpatialRequirements', ''),
            'success_criteria': data.get('SuccessCriteria', ''),
            'failure_indicators': data.get('FailureIndicators', '')
        }
    except Exception as e:
        print(f"⚠️  Could not load Excel: {e}")
        return None


def find_video(folder: str, name: str) -> Optional[str]:
    """Find video file"""
    
    if not os.path.exists(folder):
        return None
    
    for file in os.listdir(folder):
        if name.lower() in file.lower() and file.lower().endswith(('.mp4', '.avi', '.mov')):
            return os.path.join(folder, file)
    
    return None


def save_json(data: Dict, path: Path):
    """Save JSON file"""
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


# =============================================================================
# 4. FRAME EXTRACTOR
# =============================================================================

class FrameExtractor:
    """Extract key frames from video"""
    
    def __init__(self, config: Config):
        self.config = config
        self.hands = None
        
        if MP_AVAILABLE:
            mp_hands = mp.solutions.hands
            self.hands = mp_hands.Hands(
                static_image_mode=False,
                max_num_hands=2,
                min_detection_confidence=0.5
            )
    
    def extract(self, video_path: str, n_frames: int) -> List[Dict]:
        """Extract n key frames"""
        
        cap = cv2.VideoCapture(video_path)
        if not cap.isOpened():
            return []
        
        # Read all frames
        frames = []
        idx = 0
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            frames.append({'frame': frame, 'idx': idx, 'score': 0.0})
            idx += 1
        
        cap.release()
        
        if len(frames) == 0:
            return []
        
        # Calculate motion scores
        if self.hands:
            self._score_frames(frames)
        
        # Select top N frames
        sorted_frames = sorted(frames, key=lambda x: x['score'], reverse=True)
        selected = sorted_frames[:min(n_frames, len(sorted_frames))]
        selected = sorted(selected, key=lambda x: x['idx'])
        
        # Encode to base64
        result = []
        for frame_info in selected:
            b64 = self._encode(frame_info['frame'])
            if b64:
                result.append({
                    'frame_idx': frame_info['idx'],
                    'image_base64': b64
                })
        
        return result
    
    def _score_frames(self, frames: List[Dict]):
        """Score frames by motion"""
        
        prev_gray = None
        for frame_info in frames:
            frame = frame_info['frame']
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            
            if prev_gray is not None:
                diff = cv2.absdiff(prev_gray, gray)
                frame_info['score'] = np.mean(diff)
            
            prev_gray = gray
    
    def _encode(self, frame: np.ndarray) -> str:
        """Encode frame to base64"""
        
        try:
            h, w = frame.shape[:2]
            if w > self.config.IMAGE_WIDTH:
                scale = self.config.IMAGE_WIDTH / w
                frame = cv2.resize(frame, (self.config.IMAGE_WIDTH, int(h * scale)))
            
            _, buffer = cv2.imencode('.jpg', frame, 
                                    [cv2.IMWRITE_JPEG_QUALITY, self.config.JPEG_QUALITY])
            return base64.b64encode(buffer).decode('utf-8')
        except:
            return ""


# =============================================================================
# 5. LLM ANALYZER (KEY COMPONENT)
# =============================================================================

class LLMAnalyzer:
    """LLM-based analyzer with OPTIMIZED PROMPT"""
    
    def __init__(self, config: Config):
        self.config = config
        self.client = None
        self.ref_frames = None
        
        if OPENAI_AVAILABLE:
            self.client = openai.OpenAI(api_key=config.OPENAI_API_KEY)
            self._load_reference()
    
    def _load_reference(self):
        """Load reference video"""
        
        ref_path = find_video(self.config.VIDEO_FOLDER, self.config.REFERENCE_VIDEO)
        if not ref_path:
            print(f"⚠️  Reference video not found: {self.config.REFERENCE_VIDEO}")
            return
        
        extractor = FrameExtractor(self.config)
        self.ref_frames = extractor.extract(ref_path, self.config.REF_FRAMES)
        
        if self.ref_frames:
            print(f"✅ Loaded {len(self.ref_frames)} reference frames")
    
    def analyze(self, frames: List[Dict], task_data: Dict) -> Optional[Result]:
        """Analyze operation"""
        
        if not self.client or not self.ref_frames:
            return None
        
        try:
            # Prepare content
            content = self._prepare_content(frames)
            
            # Build prompt
            prompt = self._build_prompt(task_data)
            
            # Call LLM
            response = self.client.chat.completions.create(
                model=self.config.GPT_MODEL,
                temperature=self.config.GPT_TEMPERATURE,
                messages=[{
                    "role": "user",
                    "content": content + [{"type": "text", "text": prompt}]
                }]
            )
            
            # Parse response
            result = self._parse_response(response.choices[0].message.content)
            
            if result:
                # Apply decision logic
                result = self._make_decision(result)
            
            return result
            
        except Exception as e:
            print(f"❌ Analysis error: {e}")
            return None
    
    def _prepare_content(self, frames: List[Dict]) -> List[Dict]:
        """Prepare content for LLM"""
        
        content = []
        
        # Reference frames
        content.append({"type": "text", "text": "\n=== REFERENCE (STANDARD OPERATION) ===\n"})
        for i, ref in enumerate(self.ref_frames, 1):
            content.append({"type": "text", "text": f"Ref-{i}:"})
            content.append({
                "type": "image_url",
                "image_url": {
                    "url": f"data:image/jpeg;base64,{ref['image_base64']}",
                    "detail": "high"
                }
            })
        
        # Current frames
        content.append({"type": "text", "text": "\n=== CURRENT OPERATION (TO EVALUATE) ===\n"})
        for i, frame in enumerate(frames, 1):
            content.append({"type": "text", "text": f"Current-{i}:"})
            content.append({
                "type": "image_url",
                "image_url": {
                    "url": f"data:image/jpeg;base64,{frame['image_base64']}",
                    "detail": "high"
                }
            })
        
        return content
    
    def _build_prompt(self, task_data: Dict) -> str:
        """Build optimized GENERIC prompt"""
        
        prompt = f"""You are comparing two assembly operations: REFERENCE (correct) vs CURRENT (to evaluate).

TASK INFORMATION:
- Description: {task_data.get('task_description', 'N/A')}
- Success Criteria: {task_data.get('success_criteria', 'N/A')}
- Failure Indicators: {task_data.get('failure_indicators', 'N/A')}
- Safety Requirements: {task_data.get('safety_requirements', 'N/A')}
- Spatial Requirements: {task_data.get('spatial_requirements', 'N/A')}

{'='*70}
ANALYSIS METHOD: OUTCOME-FOCUSED COMPARISON
{'='*70}

STEP 1: Identify ALL differences between CURRENT and REFERENCE
- Look at every aspect: positioning, timing, hand movements, component placement, etc.

STEP 2: For EACH difference, answer these questions:
   A) Does this difference affect the FINAL OUTCOME of the task?
      - Look at the END STATE, not just the process
      - Did the task achieve the same result as reference?
   
   B) Does this difference violate the Success Criteria or trigger Failure Indicators?
      - Check against the task requirements above
   
   C) Classify the difference:
      
      CRITICAL (is_problem=True):
      - Task completely fails or fundamentally incorrect
      - Severe safety violations
      - Component in wrong location or not installed
      - Any failure indicator from task data is present
      
      MAJOR (is_problem=True):
      - Task incomplete or significantly deviated
      - Quality issues affecting function
      - Procedure violations affecting outcome
      - Partially meets criteria but has defects
      
      MINOR (is_problem=True):
      - Task complete but with small quality issues
      - Minor deviations that slightly affect result
      - Not ideal but acceptable
      
      NEGLIGIBLE (is_problem=False):
      - Different approach but same result
      - Timing/speed differences
      - Camera angle or lighting effects
      - Process variations that don't affect outcome

KEY DECISION RULES:

🎯 Rule 1: COMPARE OUTCOMES, NOT JUST PROCESSES
   - If CURRENT achieves same END STATE as REFERENCE → Likely NEGLIGIBLE
   - If END STATE is different → Likely MAJOR or CRITICAL

🎯 Rule 2: USE THE TASK CRITERIA
   - Check: Does CURRENT meet ALL Success Criteria?
   - Check: Does CURRENT show ANY Failure Indicators?
   - If yes to any Failure Indicator → Definitely a PROBLEM

🎯 Rule 3: CONSIDER SEVERITY BY IMPACT
   - Would this difference cause the assembly to FAIL in real use? → CRITICAL
   - Would this difference cause QUALITY issues? → MAJOR
   - Is this just a PROCEDURAL difference with same result? → NEGLIGIBLE

SCORING GUIDANCE:
- CURRENT ≈ REFERENCE outcome, all criteria met: 85-100
- CURRENT has minor deviations from ideal: 65-84
- CURRENT has significant issues but task somewhat done: 40-64
- CURRENT fails task or has critical problems: 0-39

CRITICAL: If both operations show similar features (e.g., both have small gaps),
that's NOT a problem - it's part of the normal acceptable range!

{'='*70}

Return JSON:
{{
    "safety_score": <0-100>,
    "spatial_score": <0-100>,
    "operation_score": <0-100>,
    "overall_score": <0-100>,
    "summary": "<what you observed overall>",
    "observed_differences": [
        {{
            "description": "<what is different>",
            "is_problem": <true/false>,
            "severity": "<CRITICAL|MAJOR|MINOR|NEGLIGIBLE>",
            "category": "<safety|spatial|operation>"
        }}
    ]
}}
"""
        
        return prompt
    
    def _parse_response(self, text: str) -> Optional[Result]:
        """Parse LLM response"""
        
        try:
            # Extract JSON
            start = text.find('{')
            end = text.rfind('}') + 1
            if start < 0 or end <= start:
                return None
            
            data = json.loads(text[start:end])
            
            # Parse differences
            differences = []
            for d in data.get('observed_differences', []):
                differences.append(Difference(
                    description=d.get('description', ''),
                    is_problem=d.get('is_problem', False),
                    severity=d.get('severity', 'NEGLIGIBLE').upper(),
                    category=d.get('category', 'operation')
                ))
            
            result = Result(
                video_name="",
                timestamp=datetime.now().isoformat(),
                safety_score=float(data.get('safety_score', 0)),
                spatial_score=float(data.get('spatial_score', 0)),
                operation_score=float(data.get('operation_score', 0)),
                overall_score=float(data.get('overall_score', 0)),
                differences=differences,
                summary=data.get('summary', '')
            )
            
            return result
            
        except Exception as e:
            print(f"❌ Parse error: {e}")
            return None
    
    def _make_decision(self, result: Result) -> Result:
        """Apply decision logic"""
        
        # Count problems
        n_critical = len(result.critical)
        n_major = len(result.major)
        
        # Rule 1: Critical problems
        if n_critical >= self.config.CRITICAL_THRESHOLD:
            result.decision = "FAIL"
            result.reason = f"Critical problem(s): {n_critical}"
            return result
        
        # Rule 2: Major problems
        if n_major >= self.config.MAJOR_THRESHOLD:
            result.decision = "FAIL"
            result.reason = f"Major problems: {n_major}"
            return result
        
        # Rule 3: Low score
        if result.overall_score < self.config.SCORE_THRESHOLD:
            result.decision = "FAIL"
            result.reason = f"Score: {result.overall_score:.0f} < {self.config.SCORE_THRESHOLD}"
            return result
        
        # Pass
        result.decision = "PASS"
        result.reason = f"Score: {result.overall_score:.0f}, Critical: {n_critical}, Major: {n_major}"
        
        return result


# =============================================================================
# 6. MAIN SYSTEM
# =============================================================================

class DetectionSystem:
    """Main detection system"""
    
    def __init__(self, config: Config):
        self.config = config
        self.extractor = FrameExtractor(config)
        self.analyzer = LLMAnalyzer(config)
        self.task_data = load_task_data(config)
    
    def process_video(self, video_name: str) -> Optional[Result]:
        """Process one video"""
        
        print(f"\n{'='*70}")
        print(f"PROCESSING: {video_name}")
        print(f"{'='*70}")
        
        # Find video
        video_path = find_video(self.config.VIDEO_FOLDER, video_name)
        if not video_path:
            print(f"❌ Video not found")
            return None
        
        # Extract frames
        print(f"📹 Extracting frames...")
        t0 = time.time()
        frames = self.extractor.extract(video_path, self.config.KEY_FRAMES)
        print(f"   ✅ Extracted {len(frames)} frames ({time.time()-t0:.1f}s)")
        
        if not frames:
            return None
        
        # Save frames
        self._save_frames(video_name, frames)
        
        # Analyze
        print(f"🤖 Analyzing...")
        t0 = time.time()
        result = self.analyzer.analyze(frames, self.task_data or {})
        print(f"   ✅ Analysis complete ({time.time()-t0:.1f}s)")
        
        if not result:
            return None
        
        # Add metadata
        result.video_name = video_name
        result.ground_truth = GROUND_TRUTH.get(video_name, "UNKNOWN")
        result.correct = (result.ground_truth == "OK" and result.decision == "PASS") or \
                        (result.ground_truth == "KO" and result.decision == "FAIL")
        
        # Display and save
        self._display(result)
        self._save_result(result)
        
        return result
    
    def process_batch(self, video_names: List[str]) -> List[Result]:
        """Process multiple videos"""
        
        results = []
        for i, name in enumerate(video_names, 1):
            print(f"\n[{i}/{len(video_names)}] {name}")
            result = self.process_video(name)
            if result:
                results.append(result)
            time.sleep(0.5)
        
        self._display_summary(results)
        return results
    
    def _save_frames(self, video_name: str, frames: List[Dict]):
        """Save frames to disk"""
        
        output_dir = Path(self.config.RESULTS_FOLDER) / "frames" / video_name
        output_dir.mkdir(parents=True, exist_ok=True)
        
        for frame in frames:
            idx = frame['frame_idx']
            data = base64.b64decode(frame['image_base64'])
            
            with open(output_dir / f"frame_{idx:04d}.jpg", 'wb') as f:
                f.write(data)
    
    def _save_result(self, result: Result):
        """Save result to JSON"""
        
        output_path = Path(self.config.RESULTS_FOLDER) / "reports" / f"{result.video_name}.json"
        
        data = {
            'video_name': result.video_name,
            'timestamp': result.timestamp,
            'scores': {
                'safety': result.safety_score,
                'spatial': result.spatial_score,
                'operation': result.operation_score,
                'overall': result.overall_score
            },
            'differences': {
                'total': len(result.differences),
                'critical': [{'desc': d.description, 'cat': d.category} for d in result.critical],
                'major': [{'desc': d.description, 'cat': d.category} for d in result.major],
                'minor': [{'desc': d.description, 'cat': d.category} for d in result.minor],
                'negligible': [{'desc': d.description, 'cat': d.category} for d in result.negligible]
            },
            'decision': {
                'result': result.decision,
                'reason': result.reason
            },
            'validation': {
                'ground_truth': result.ground_truth,
                'correct': result.correct
            },
            'summary': result.summary
        }
        
        save_json(data, output_path)
        print(f"💾 Saved: {output_path.name}")
    
    def _display(self, result: Result):
        """Display result"""
        
        print(f"\n{'='*70}")
        print(f"RESULT: {result.video_name}")
        print(f"{'='*70}")
        print(f"📊 Scores: Safety={result.safety_score:.0f}, Spatial={result.spatial_score:.0f}, "
              f"Operation={result.operation_score:.0f}, Overall={result.overall_score:.0f}")
        print(f"🔍 Differences: Total={len(result.differences)}, Critical={len(result.critical)}, "
              f"Major={len(result.major)}, Minor={len(result.minor)}, Negligible={len(result.negligible)}")
        
        if result.critical:
            print(f"\n🔴 CRITICAL:")
            for d in result.critical:
                print(f"   - {d.description}")
        
        if result.major:
            print(f"\n🟠 MAJOR:")
            for d in result.major:
                print(f"   - {d.description}")
        
        print(f"\n💡 Decision: {result.decision}")
        print(f"📝 Reason: {result.reason}")
        print(f"✅ Validation: GT={result.ground_truth}, {'✅ CORRECT' if result.correct else '❌ WRONG'}")
        print(f"{'='*70}")
    
    def _display_summary(self, results: List[Result]):
        """Display batch summary"""
        
        if not results:
            return
        
        print(f"\n{'='*70}")
        print(f"BATCH SUMMARY")
        print(f"{'='*70}")
        
        total = len(results)
        correct = sum(1 for r in results if r.correct)
        accuracy = correct / total * 100 if total > 0 else 0
        
        print(f"\nTotal: {total}")
        print(f"Correct: {correct}/{total} ({accuracy:.1f}%)")
        
        print(f"\n{'Video':<10} {'GT':<4} {'Pred':<6} {'Score':<6} {'Crit':<5} {'Major':<6} {'Result'}")
        print(f"{'-'*70}")
        
        for r in results:
            pred = "OK" if r.decision == "PASS" else "KO"
            status = "✅" if r.correct else "❌"
            print(f"{r.video_name:<10} {r.ground_truth:<4} {pred:<6} "
                  f"{r.overall_score:<6.0f} {len(r.critical):<5} {len(r.major):<6} {status}")
        
        print(f"{'-'*70}")
        
        # Save summary
        summary_path = Path(self.config.RESULTS_FOLDER) / "reports" / "summary.json"
        save_json({
            'timestamp': datetime.now().isoformat(),
            'total': total,
            'correct': correct,
            'accuracy': accuracy,
            'results': [
                {
                    'video': r.video_name,
                    'gt': r.ground_truth,
                    'decision': r.decision,
                    'score': r.overall_score,
                    'critical': len(r.critical),
                    'major': len(r.major),
                    'correct': r.correct
                }
                for r in results
            ]
        }, summary_path)
        print(f"\n💾 Summary saved: {summary_path.name}")


# =============================================================================
# 7. MAIN
# =============================================================================

def main():
    """Main entry"""
    
    print(f"\n{'='*70}")
    print(f"ASSEMBLY DETECTION SYSTEM V6.0 (Simplified)")
    print(f"{'='*70}\n")
    
    config = Config()
    
    print(f"📋 Config:")
    print(f"   Model: {config.GPT_MODEL}")
    print(f"   Reference: {config.REFERENCE_VIDEO}")
    print(f"   Thresholds: Critical≥{config.CRITICAL_THRESHOLD}, Major≥{config.MAJOR_THRESHOLD}, Score<{config.SCORE_THRESHOLD}")
    
    system = DetectionSystem(config)
    
    # Get videos to process
    videos = ["OK1", "OK2", "OK3", "OK4", "OK5", "KO1", "KO2", "KO3", "KO4", "KO5"]
    
    choice = input(f"\nProcess all {len(videos)} videos? (y/n): ").strip().lower()
    
    if choice == 'y':
        system.process_batch(videos)
    else:
        name = input("Enter video name (e.g., OK1): ").strip()
        system.process_video(name)
    
    print(f"\n{'='*70}")
    print(f"DONE")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    main()